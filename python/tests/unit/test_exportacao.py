"""Exportação CSV/XLSX (AC-036): cabeçalho datado com solicitante, sem injeção de fórmula."""

import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from app.utils.exportacao import gerar_csv, gerar_xlsx

QUANDO = datetime(2026, 3, 20, 14, 31, 7)
COLUNAS = ["id", "nome", "valor"]
LINHAS = [[1, "Notebook", Decimal("1234.50")], [2, '=HYPERLINK("http://x")', Decimal("0.10")]]


def test_csv_traz_data_hora_e_usuario_no_cabecalho():
    texto = gerar_csv("Inventário", QUANDO, "admin", COLUNAS, LINHAS).decode("utf-8")

    linhas = texto.splitlines()
    assert linhas[0] == "Relatório: Inventário"
    assert linhas[1] == "Gerado em: 2026-03-20 14:31:07"
    assert linhas[2] == "Solicitante: admin"
    assert linhas[4] == "id,nome,valor"
    assert linhas[5] == "1,Notebook,1234.50"  # Decimal sai como texto, sem float


def test_csv_neutraliza_formula_em_texto_mas_nao_em_numero():
    texto = gerar_csv("R", QUANDO, "admin", COLUNAS, [[1, "+cmd", Decimal("-5.00")]]).decode()

    assert "'+cmd" in texto  # texto de usuário que parece fórmula
    assert "-5.00" in texto and "'-5.00" not in texto  # número negativo não é tocado


def test_csv_serializa_data_e_vazio_e_rodape():
    texto = gerar_csv(
        "R", QUANDO, "admin", ["d", "x"], [[date(2025, 1, 10), None]], rodape=[["TOTAL", 1]]
    ).decode()

    assert "2025-01-10," in texto
    assert texto.splitlines()[-1] == "TOTAL,1"


def test_xlsx_tem_o_mesmo_cabecalho_e_linhas():
    conteudo = gerar_xlsx("Inventário", QUANDO, "admin", COLUNAS, LINHAS, rodape=[["TOTAL", 2]])

    folha = load_workbook(io.BytesIO(conteudo)).active
    valores = [[c.value for c in linha] for linha in folha.iter_rows()]
    assert valores[0][0] == "Relatório: Inventário"
    assert valores[1][0] == "Gerado em: 2026-03-20 14:31:07"
    assert valores[2][0] == "Solicitante: admin"
    assert valores[4] == COLUNAS
    assert valores[5][:2] == [1, "Notebook"]
    assert valores[6][1] == '\'=HYPERLINK("http://x")'
    assert valores[-1][:2] == ["TOTAL", 2]
