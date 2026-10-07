"""Relatório de inventário (FR-006): AC-031, AC-034 a AC-038, e o de conformidade (FR-007).

Depreciação com valor 6.000,00 e vida útil de 60 meses: cada mês completo deprecia 100,00, então
o esperado é a conta `100 × meses` escrita no teste, nunca copiada do retorno.
"""

import csv
import io
from calendar import monthrange
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.models.ativo import Ativo
from app.models.categoria import Categoria
from app.utils import datas


def _headers(token):
    return {"Authorization": f"Bearer {token}"}


def _meses_atras(meses: int) -> str:
    hoje = datas.hoje()
    ano, mes = divmod(hoje.year * 12 + hoje.month - 1 - meses, 12)
    return date(ano, mes + 1, min(hoje.day, monthrange(ano, mes + 1)[1])).isoformat()


def _inventario(client, token, **params):
    resposta = client.get("/api/v1/relatorios/inventario", headers=_headers(token), params=params)
    assert resposta.status_code == 200, resposta.text
    return resposta.json()


def _ids(corpo):
    return sorted(i["id"] for i in corpo["itens"])


def _baixar(client, token, ativo_id):
    resposta = client.post(
        f"/api/v1/ativos/{ativo_id}/baixa",
        headers=_headers(token),
        json={"motivo": "DEFEITO", "data_baixa": "2025-07-10", "destinacao": "DESCARTE"},
    )
    assert resposta.status_code == 201, resposta.text


def test_ac031_baixado_fica_fora_do_inventario_ativo_e_volta_com_o_filtro(
    client, token_admin, novo_ativo
):
    operando, baixado = novo_ativo(), novo_ativo()
    _baixar(client, token_admin, baixado)

    assert _ids(_inventario(client, token_admin)) == [operando]
    com_filtro = _inventario(client, token_admin, status="BAIXADO")
    assert _ids(com_filtro) == [baixado]
    assert com_filtro["itens"][0]["status"] == "BAIXADO"


def test_ac034_filtros_de_status_categoria_e_responsavel_valem_em_conjuncao(
    client, db, token_admin, novo_ativo, categoria, responsavel, setor
):
    outra = Categoria(nome="Monitor", vida_util_meses=60, tipo_aplicavel="HARDWARE")
    db.add(outra)
    db.commit()
    alvo = novo_ativo()
    em_manutencao = novo_ativo()
    outra_categoria = novo_ativo(categoria_id=outra.id)
    sem_responsavel = novo_ativo()
    for ativo in (alvo, em_manutencao, outra_categoria):
        client.post(
            f"/api/v1/ativos/{ativo}/responsavel",
            headers=_headers(token_admin),
            json={
                "responsavel_id": responsavel.id,
                "setor_id": setor.id,
                "data_inicio": "2025-02-01",
            },
        )
    db.query(Ativo).filter(Ativo.id == em_manutencao).update({"status": "EM_MANUTENCAO"})
    db.commit()

    um_filtro = _inventario(client, token_admin, categoria_id=categoria.id)
    todos = _inventario(
        client,
        token_admin,
        status="ATIVO",
        categoria_id=categoria.id,
        responsavel_id=responsavel.id,
    )

    assert _ids(um_filtro) == sorted([alvo, em_manutencao, sem_responsavel])
    assert _ids(todos) == [alvo]
    assert todos["itens"][0]["responsavel"] == "Carlos Menezes"
    assert todos["itens"][0]["setor"] == "Infraestrutura"


def test_filtros_multiplos_de_tipo_e_fornecedor(client, token_admin, novo_ativo, fornecedor):
    hardware, software = novo_ativo(), novo_ativo(tipo="SOFTWARE")

    assert _ids(_inventario(client, token_admin, tipo="SOFTWARE")) == [software]
    assert _ids(_inventario(client, token_admin, tipo=["HARDWARE", "SOFTWARE"])) == sorted(
        [hardware, software]
    )
    assert _ids(_inventario(client, token_admin, fornecedor_id=fornecedor.id)) == sorted(
        [hardware, software]
    )
    assert _inventario(client, token_admin, fornecedor_id=fornecedor.id + 99)["itens"] == []


def test_ac035_faixa_de_valor_depreciado_so_traz_ativos_na_faixa(client, token_admin, novo_ativo):
    por_meses = {
        meses: novo_ativo(valor="6000.00", data_aquisicao=_meses_atras(meses))
        for meses in (5, 12, 20, 30, 31)
    }  # depreciado = 100 × meses: 500, 1.200, 2.000, 3.000, 3.100

    corpo = _inventario(
        client, token_admin, valor_depreciado_min="1000.00", valor_depreciado_max="3000.00"
    )

    assert _ids(corpo) == sorted([por_meses[12], por_meses[20], por_meses[30]])
    assert {i["id"]: i["depreciacao_acumulada"] for i in corpo["itens"]} == {
        por_meses[12]: "1200.00",
        por_meses[20]: "2000.00",
        por_meses[30]: "3000.00",
    }


def test_faixa_de_percentual_depreciado(client, token_admin, novo_ativo):
    por_meses = {
        m: novo_ativo(valor="6000.00", data_aquisicao=_meses_atras(m)) for m in (3, 6, 30, 54)
    }

    corpo = _inventario(
        client, token_admin, percentual_depreciado_min="10", percentual_depreciado_max="50"
    )  # 3 meses = 5%; 10% = 6 meses; 50% = 30 meses; 54 meses = 90%

    assert _ids(corpo) == sorted([por_meses[6], por_meses[30]])


def test_periodo_de_aquisicao(client, token_admin, novo_ativo):
    antigo = novo_ativo(data_aquisicao="2024-01-10")
    novo = novo_ativo(data_aquisicao="2025-06-10")

    assert _ids(_inventario(client, token_admin, aquisicao_de="2025-01-01")) == [novo]
    assert _ids(_inventario(client, token_admin, aquisicao_ate="2024-12-31")) == [antigo]


def test_ac038_fim_de_vida_util_traz_80_por_cento_ou_mais(client, token_admin, novo_ativo):
    por_meses = {
        m: novo_ativo(valor="6000.00", data_aquisicao=_meses_atras(m)) for m in (47, 48, 60)
    }
    # vida útil de 60 meses: 47 = 78,33% · 48 = 80,00% · 60 = 100,00%

    corpo = _inventario(client, token_admin, fim_vida_util="true")

    assert _ids(corpo) == sorted([por_meses[48], por_meses[60]])


def test_baixado_usa_o_residual_congelado_na_data_da_baixa(client, token_admin, novo_ativo):
    ativo = novo_ativo(valor="6000.00", data_aquisicao="2025-01-10")  # baixa em 2025-07-10: 6 meses
    _baixar(client, token_admin, ativo)

    (linha,) = _inventario(client, token_admin, status="BAIXADO")["itens"]

    assert linha["valor_residual"] == "5400.00"  # 6.000 − 100 × 6
    assert linha["depreciacao_acumulada"] == "600.00"
    assert linha["percentual_depreciado"] == "10.00"


def test_totalizadores_e_paginacao_do_json(client, token_admin, novo_ativo):
    for _ in range(3):
        novo_ativo(valor="6000.00", data_aquisicao=_meses_atras(10))  # residual 6.000 − 1.000

    corpo = _inventario(client, token_admin, tamanho=2, pagina=2)

    assert (corpo["pagina"], corpo["tamanho"], corpo["total"], corpo["total_paginas"]) == (
        2,
        2,
        3,
        2,
    )
    assert len(corpo["itens"]) == 1
    assert corpo["totais"] == {  # sobre o resultado todo, não só a página
        "quantidade": 3,
        "valor_compra": "18000.00",
        "valor_residual": "15000.00",
    }


def test_chave_do_software_sai_mascarada_no_relatorio(client, token_admin, novo_ativo):
    novo_ativo(tipo="SOFTWARE")

    (linha,) = _inventario(client, token_admin)["itens"]

    assert linha["identificador"].startswith("****-****-")
    assert "K-T1" not in linha["identificador"]


def test_ac036_csv_traz_data_hora_e_usuario_solicitante(client, token_admin, novo_ativo):
    novo_ativo()

    resposta = client.get(
        "/api/v1/relatorios/inventario", headers=_headers(token_admin), params={"formato": "csv"}
    )

    assert resposta.status_code == 200
    assert resposta.headers["content-type"].startswith("text/csv")
    assert 'filename="inventario.csv"' in resposta.headers["content-disposition"]
    linhas = resposta.text.splitlines()
    assert linhas[0] == "Relatório: Inventário de ativos"
    assert linhas[1].startswith(f"Gerado em: {date.today().isoformat()} ")  # data e hora
    assert linhas[2] == "Solicitante: admin_teste"
    assert linhas[4].startswith("id,nome,tipo,categoria,identificador,responsavel,setor,status")


def test_ac037_exportacao_traz_todos_os_500_ativos_sem_truncar_por_paginacao(
    client, db, token_admin, categoria, fornecedor
):
    db.add_all(
        Ativo(
            nome=f"Lote {n:03d}",
            tipo="HARDWARE",
            categoria_id=categoria.id,
            fornecedor_id=fornecedor.id,
            numero_serie=f"SN-LOTE-{n}",
            data_aquisicao=date(2025, 1, 10),
            valor_compra=Decimal("1000.00"),
            vida_util_meses=60,
            status="ATIVO",
            data_source="manual",
        )
        for n in range(500)
    )
    db.commit()

    resposta = client.get(
        "/api/v1/relatorios/inventario",
        headers=_headers(token_admin),
        params={"formato": "csv", "tamanho": 20, "pagina": 1},  # a paginação é ignorada
    )

    linhas = list(csv.reader(io.StringIO(resposta.text)))
    dados = [c for c in linhas[5:] if c and c[0].isdigit()]
    assert len(dados) == 500
    assert linhas[-1][0:2] == ["TOTAL", "500"]
    assert linhas[-1][9] == "500000.00"  # 500 × 1.000,00


def test_exportacao_xlsx(client, token_admin, novo_ativo):
    ativo = novo_ativo()

    resposta = client.get(
        "/api/v1/relatorios/inventario", headers=_headers(token_admin), params={"formato": "xlsx"}
    )

    assert resposta.status_code == 200
    assert "spreadsheetml" in resposta.headers["content-type"]
    folha = load_workbook(io.BytesIO(resposta.content)).active
    valores = [[c.value for c in linha] for linha in folha.iter_rows()]
    assert valores[2][0] == "Solicitante: admin_teste"
    assert valores[4][0] == "id"
    assert valores[5][0] == ativo


def test_formato_desconhecido_e_422(client, token_admin):
    resposta = client.get(
        "/api/v1/relatorios/inventario", headers=_headers(token_admin), params={"formato": "pdf"}
    )

    assert resposta.status_code == 422


def test_relatorio_de_conformidade_datado_em_json_csv_e_xlsx(
    client, token_admin, nova_licenca, novo_ativo
):
    nova_licenca(inicio="2019-01-01", expiracao="2020-01-01")
    novo_ativo()

    json_ = client.get("/api/v1/relatorios/conformidade", headers=_headers(token_admin)).json()
    csv_ = client.get(
        "/api/v1/relatorios/conformidade", headers=_headers(token_admin), params={"formato": "csv"}
    )
    xlsx = client.get(
        "/api/v1/relatorios/conformidade", headers=_headers(token_admin), params={"formato": "xlsx"}
    )

    assert json_["total"] == 2
    linhas = csv_.text.splitlines()
    assert linhas[0] == f"Relatório: Conformidade em {datas.hoje().isoformat()}"
    assert linhas[2] == "Solicitante: admin_teste"
    assert linhas[4] == "codigo,severidade,regra,mensagem,entidade,entidade_id,recurso"
    assert {linha.split(",")[0] for linha in linhas[5:]} == {"CP-01", "CP-04"}
    assert load_workbook(io.BytesIO(xlsx.content)).active["A1"].value.startswith("Relatório")
